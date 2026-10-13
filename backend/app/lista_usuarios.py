# Lista de usuarios para el alta masiva: lee un Excel (.xlsx) o CSV con las columnas
# nombre, correo, cargo (opcional) y rol, y revisa cada fila antes de crear nada.
import csv
import io
import re
import unicodedata

from openpyxl import load_workbook

from .errores import ErrorApi
from .seguridad import problema_con_correo

MAX_BYTES = 2 * 1048576
MAX_FILAS = 500

# Encabezado (normalizado) -> campo. Se aceptan los nombres habituales de una planilla.
COLUMNAS = {
    'nombre': 'nombre_completo', 'nombre completo': 'nombre_completo', 'nombres y apellidos': 'nombre_completo',
    'apellidos y nombres': 'nombre_completo', 'nombres': 'nombre_completo',
    'correo': 'correo', 'correo electronico': 'correo', 'email': 'correo', 'e mail': 'correo', 'mail': 'correo',
    'cargo': 'cargo', 'puesto': 'cargo',
    'rol': 'rol', 'perfil': 'rol',
}
ROLES = {
    'admin': 'ADMIN', 'administrador': 'ADMIN', 'administradora': 'ADMIN',
    'jefe': 'JEFE', 'jefa': 'JEFE',
    'abogado': 'ABOGADO', 'abogada': 'ABOGADO', 'profesional': 'ABOGADO',
    'secretaria': 'SECRETARIA', 'secretario': 'SECRETARIA',
}


def _norm(v):
    v = unicodedata.normalize('NFD', str(v or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[\s_\-]+', ' ', v).strip()


def _texto(v):
    return re.sub(r'\s+', ' ', str(v).strip()) if v is not None else ''


def _filas_del_archivo(nombre, contenido):
    nombre = (nombre or '').lower()
    if nombre.endswith('.xlsx') or contenido[:2] == b'PK':
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise ErrorApi('No se pudo leer el Excel. Guárdelo como .xlsx e intente de nuevo.', 422)
        try:
            return [list(f) for f in libro.worksheets[0].iter_rows(values_only=True)]
        finally:
            libro.close()
    if nombre.endswith(('.csv', '.txt')):
        try:
            texto = contenido.decode('utf-8-sig')
        except UnicodeDecodeError:
            texto = contenido.decode('latin-1')  # CSV guardado desde Excel en Windows
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=';,\t')
        except csv.Error:
            dialecto = csv.excel
        return list(csv.reader(io.StringIO(texto), dialecto))
    raise ErrorApi('Suba un archivo Excel (.xlsx) o CSV.', 422, {'archivo': 'Formato no admitido.'})


def leer(nombre, contenido):
    """Filas con los campos reconocidos: [{fila, nombre_completo, correo, cargo, rol_original}]."""
    if not contenido:
        raise ErrorApi('Adjunte la lista de usuarios.', 422, {'archivo': 'Adjunte un archivo.'})
    if len(contenido) > MAX_BYTES:
        raise ErrorApi('El archivo supera 2 MB.', 422, {'archivo': 'Máximo 2 MB.'})
    filas = _filas_del_archivo(nombre, contenido)
    # El encabezado es la primera de las 10 primeras filas que tenga la columna del correo.
    for i, fila in enumerate(filas[:10]):
        columnas = {j: COLUMNAS[_norm(c)] for j, c in enumerate(fila) if _norm(c) in COLUMNAS}
        if 'correo' in columnas.values():
            break
    else:
        raise ErrorApi('No se encontró la fila de títulos: debe tener al menos las columnas Nombre, Correo y Rol.', 422)
    faltan = {'nombre_completo', 'correo', 'rol'} - set(columnas.values())
    if faltan:
        nombres = {'nombre_completo': 'Nombre', 'correo': 'Correo', 'rol': 'Rol'}
        raise ErrorApi(f"Faltan columnas: {', '.join(nombres[f] for f in sorted(faltan))}.", 422)

    resultado = []
    for n, fila in enumerate(filas[i + 1:], start=i + 2):
        datos = {campo: _texto(fila[j]) if j < len(fila) else '' for j, campo in columnas.items()}
        if not any(datos.values()):
            continue  # fila vacía
        resultado.append({'fila': n, 'nombre_completo': datos.get('nombre_completo', ''),
                          'correo': datos.get('correo', '').lower(), 'cargo': datos.get('cargo', ''),
                          'rol_original': datos.get('rol', '')})
        if len(resultado) > MAX_FILAS:
            raise ErrorApi(f'La lista tiene más de {MAX_FILAS} usuarios. Divídala en varias partes.', 422)
    if not resultado:
        raise ErrorApi('La lista no tiene usuarios debajo de la fila de títulos.', 422)
    return resultado


def revisar(con, filas):
    """Marca en cada fila el rol reconocido y el problema que impide crearla (error) o None."""
    vistos = set()
    for f in filas:
        f['rol'] = ROLES.get(_norm(f.pop('rol_original')))
        if not f['nombre_completo']:
            f['error'] = 'Falta el nombre.'
        elif len(f['nombre_completo']) > 150:
            f['error'] = 'El nombre supera 150 caracteres.'
        elif f['correo'] in vistos:
            f['error'] = 'El correo está repetido en la lista.'
        elif problema := problema_con_correo(con, f['correo']):
            f['error'] = problema
        elif not f['rol']:
            f['error'] = 'Rol no reconocido: use Administrador, Jefe, Abogado o Secretaria.'
        else:
            f['error'] = None
        f['cargo'] = f['cargo'][:120]
        vistos.add(f['correo'])
    return filas
